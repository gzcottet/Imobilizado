import os
import glob
import warnings
import pandas as pd
from datetime import datetime, timedelta 
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
import unicodedata
import re
import sys

warnings.filterwarnings("ignore", category=pd.errors.SettingWithCopyWarning)

if getattr(sys, 'frozen', False):
    BASE_DIR = os.path.dirname(sys.executable) 
else:
    BASE_DIR = os.path.abspath(os.path.dirname(__file__)) 

INPUT_DIR = os.path.join(BASE_DIR, "data", "input")
OUTPUT_DIR = os.path.join(BASE_DIR, "data", "output")

os.makedirs(INPUT_DIR, exist_ok=True)
os.makedirs(OUTPUT_DIR, exist_ok=True)

arquivos_xlsx = sorted(glob.glob(os.path.join(INPUT_DIR, "*.xlsx")))
if not arquivos_xlsx:
    print("\nERRO CRÍTICO: Nenhum arquivo .xlsx encontrado!")
    print(f"   Caminho esperado: {INPUT_DIR}")
    sys.exit(1)

INPUT_FILE = arquivos_xlsx[0]
print(f"Usando arquivo de entrada: {INPUT_FILE}")

SHEET_NAME = "Imobilizado"
colunas_para_ler = list(range(0, 31))

novo_cabecalho = [
    "código da conta", "Plaqueta", "Descrição", "Local", "Número de documento",
    "Fornecedor", "Quantidade", "Data aquisição", "Valor de aquisição", "Valor depreciado",
    "Vida útil (meses)", "Taxa de depreciação (anual)", "Função do bem na atividade do estabelecimento",
    "Complemento do local", "Código anterior do item", "Depreciação acelerada",
    "Data início depreciação", "Emissão própria", "Número", "Data", "Modelo", "Série",
    "Espécie", "Sequência na NF", "Emitente", "Chave da NF-e", "Código do produto na NF"
]

def choose_series_strict(df, candidate_names, field_label, required=False):
    normalized_cols = {str(col).strip().lower(): col for col in df.columns}
    for name in candidate_names:
        norm_name = name.strip().lower()
        if norm_name in normalized_cols:
            return df[normalized_cols[norm_name]]
    if required:
        raise ValueError(f"Campo obrigatório '{field_label}' não encontrado.")
    return pd.Series([pd.NA] * len(df), dtype='object')

def to_date(series):
    return pd.to_datetime(series, errors="coerce", dayfirst=True)

def to_number(series):
    s = series.astype(str).fillna("")
    s = s.str.replace(r"[^\d,.\-]", "", regex=True).str.replace(",", ".")
    return pd.to_numeric(s.replace("", pd.NA), errors="coerce")

def calcular_meses_entre(s_inicio, s_fim):
    start = to_date(s_inicio)
    end = to_date(s_fim)
    def meses(row):
        i, f = row
        if pd.isna(i) or pd.isna(f):
            return pd.NA
        return (f.year - i.year) * 12 + (f.month - i.month)
    return pd.Series([meses(x) for x in zip(start, end)], index=s_inicio.index)

def normalize_str(s):
    if pd.isna(s):
        return ""
    s = str(s).strip()
    s = unicodedata.normalize('NFKD', s)
    s = ''.join(ch for ch in s if not unicodedata.combining(ch))
    s = s.lower()
    s = re.sub(r'\s+', ' ', s)
    return s

def corrigir_data(dt):
    if pd.isna(dt):
        return dt
    if dt.year >= 2050:
        try:
            return dt.replace(year=2025)
        except ValueError:
            return datetime(2025, 2, 28)
    return dt 

print("\nLendo planilha de origem...")
df = pd.read_excel(INPUT_FILE, sheet_name=SHEET_NAME, header=1, usecols=colunas_para_ler, engine="openpyxl")
df.columns = df.columns.str.strip()
print(f"{len(df)} registros lidos. Colunas detectadas: {len(df.columns)}")

df_saida = pd.DataFrame(index=df.index, columns=novo_cabecalho)

df_saida["código da conta"] = choose_series_strict(df, ["Descrição Conta Contábil","Código Conta Contábil","Descricao Conta Contabil"], "código da conta")
df_saida["Plaqueta"] = ""
df_saida["Descrição"] = choose_series_strict(df, ["Descrição do Bem","Descricao do Bem","Bem","Descrição"], "Descrição")

df_saida["Local"] = df.iloc[:, 0].fillna(0).astype(int)

df_saida["Número de documento"] = choose_series_strict(df, ["Número de documento","Número Documento","NúmeroDocumento","Nota Fiscal","NotaFiscal"], "Número de documento")
df_saida["Fornecedor"] = "84.938.430/0001-49"
df_saida["Quantidade"] = "1"
df_saida["Data aquisição"] = to_date(choose_series_strict(df, ["Data Aquisição","Data_Aquisição","Data Aquisição"], "Data aquisição"))
df_saida["Valor de aquisição"] = to_number(choose_series_strict(df, [
    "Valor Aquisição","Valor da Aquisição","ValorAquisição","Valor_ Aquisição","Valor Aquisição ",
    "Valor de Aquisição","Valor Aquisicao","Valor de Aquisicao"], "Valor de aquisição", required=True))
df_saida["Valor depreciado"] = to_number(choose_series_strict(df, ["ValPU Total Depreciação","ValPU Total","Valpu Total","Valor Final Depreciação","Valor Final Depreciacao"], "Valor depreciado"))

serie_inicio = choose_series_strict(df, ["Data Início Depreciação","Data Inicio Depreciação","Data_Início_Depreciação"], "Data Início Depreciação (para Vida útil)")
serie_fim = choose_series_strict(df, ["Data Final Depreciação","Data Final","DataFinalDepreciacao","Data Final Depreciação"], "Data Final Depreciação (para Vida útil)")
df_saida["Vida útil (meses)"] = calcular_meses_entre(serie_inicio, serie_fim)

df_saida["Taxa de depreciação (anual)"] = choose_series_strict(df, ["Taxa Anual Depreciação","Taxa Anual Depreciacao","Taxa Anual"], "Taxa de depreciação (anual)")

df_saida["Taxa de depreciação (anual)"] = df_saida["Taxa de depreciação (anual)"].replace({0.1:10, 0.2:20, 0.3:30, 0.04:4, 0.25:25, 0.0625:6.25, 0.039999998:4, 0.076923077:8, 0.05:5})

df_saida["Função do bem na atividade do estabelecimento"] = choose_series_strict(df, ["Categoria","Função do bem","Função","Classe","Tipo de Bem"], "Função do bem na atividade do estabelecimento")
df_saida["Complemento do local"] = choose_series_strict(df, ["Complemento do local","Complemento"], "Complemento do local")
df_saida["Código anterior do item"] = ""
df_saida["Depreciação acelerada"] = ""
df_saida["Data início depreciação"] = to_date(serie_inicio)
df_saida["Emissão própria"] = ""
df_saida["Número"] = choose_series_strict(df, ["Nota Fiscal","NúmeroNf","Número Nf","Numero NF","NumeroNf"], "Número")
df_saida["Data"] = to_date(serie_inicio).apply(corrigir_data)
for col in ["Data aquisição", "Data início depreciação", "Data"]:
    df_saida[col] = df_saida[col].apply(corrigir_data)

mask = (df_saida["Data início depreciação"].notna() & df_saida["Data aquisição"].notna() & (df_saida["Data início depreciação"] <= df_saida["Data aquisição"]))
df_saida.loc[mask, "Data início depreciação"] = df_saida.loc[mask, "Data aquisição"] + timedelta(days=1)

df_saida["Modelo"] = ""
df_saida["Série"] = ""
df_saida["Espécie"] = ""
df_saida["Sequência na NF"] = ""
df_saida["Emitente"] = ""
df_saida["Chave da NF-e"] = ""
df_saida["Código do produto na NF"] = ""

mapa_codigos = {
    "Máquinas E Equipamentos": "1230103001",
    "Benfeitorias": "1230102001",
    "Móveis E Utensilios": "1230104001",
    "Equipamentos De Informática": "1230105001",
    "Veículos": "1230106",
    "Veículos Da Empresa": "1230106003",
    "Intangível": "1240101001",
    "Programas De Computador": "1240101001",
    "Veículos Em Demonstração": "1230106001",
    "Construções Em Andamento": "1230102002",
    "Marcas e Patentes": "1240101002",
    "Terrenos": "1230101001"
}
df_saida["Função do bem na atividade do estabelecimento"] = df_saida["Função do bem na atividade do estabelecimento"].astype(str).str.strip()
df_saida["func_norm"] = df_saida["Função do bem na atividade do estabelecimento"].apply(normalize_str)
normalized_mapa_codigos = { normalize_str(k): v for k, v in mapa_codigos.items() }
mask_missing_codigo = df_saida["código da conta"].isna() | (df_saida["código da conta"].astype(str).str.strip() == "")
candidatos = df_saida.loc[mask_missing_codigo, "func_norm"].map(normalized_mapa_codigos)
df_saida.loc[mask_missing_codigo, "código da conta"] = candidatos.fillna(df_saida.loc[mask_missing_codigo, "código da conta"])
df_saida = df_saida.drop(columns=["func_norm"])

df_saida = df_saida[df_saida["Data aquisição"].apply(lambda x: pd.isna(x) or x.year >= 2000)].copy()
colunas_essenciais = ["Descrição", "Valor de aquisição", "Data aquisição"]
df_saida = df_saida.dropna(subset=colunas_essenciais)
mask = (df_saida["Data início depreciação"].notna() & df_saida["Data aquisição"].notna() & (df_saida["Data início depreciação"] <= df_saida["Data aquisição"]))
df_saida.loc[mask, "Data início depreciação"] = df_saida.loc[mask, "Data aquisição"] + timedelta(days=1)

locais_unicos = df_saida["Local"].unique()
for local in locais_unicos:
    df_local = df_saida[df_saida["Local"] == local].copy()
    output_file_local = os.path.join(OUTPUT_DIR, f"imobilizado_{local}.xlsx")
    print(f"\nExportando Local '{local}' → {output_file_local}")
    
    with pd.ExcelWriter(output_file_local, engine='openpyxl', datetime_format='DD/MM/YYYY') as writer:
        df_local.to_excel(writer, sheet_name="Imobilizado_Reestruturado", index=False)
    
    wb = load_workbook(output_file_local)
    ws = wb["Imobilizado_Reestruturado"]
    
    for col in ws.columns:
        max_length = max(len(str(cell.value)) if cell.value else 0 for cell in col)
        ws.column_dimensions[get_column_letter(col[0].column)].width = min(max_length + 2, 50)
    
    data_cols = ["Data aquisição", "Data início depreciação", "Data"]
 
    numero_cols = ["Valor de aquisição", "Valor depreciado", "Número"] 
    
    geral_cols = ["código da conta"] 

    col_indices = {cell.value: idx+1 for idx, cell in enumerate(ws[1])}

    for col_name in data_cols:
        if col_name in col_indices:
            col_idx = col_indices[col_name]
            for row in ws.iter_rows(min_row=2, min_col=col_idx, max_col=col_idx):
                cell = row[0]
                if cell.value is not None:
                    if isinstance(cell.value, datetime):
                        cell.value = cell.value.date()
                        cell.number_format = "DD/MM/YYYY"
                    elif isinstance(cell.value, str):
                        try:
                            dt = pd.to_datetime(cell.value, dayfirst=True, errors='raise')
                            cell.value = dt.date()
                            cell.number_format = "DD/MM/YYYY"
                        except Exception:
                            pass

    for col_name in numero_cols:
        if col_name in col_indices:
            col_idx = col_indices[col_name]
            for row in ws.iter_rows(min_row=2, min_col=col_idx, max_col=col_idx):
                cell = row[0]
                if cell.value is not None and cell.value != "":
                    try:
                        cell.value = float(cell.value)
                        cell.number_format = "#,##0.00"
                    except (ValueError, TypeError):
                        pass

  
    for col_name in geral_cols:
        if col_name in col_indices:
            col_idx = col_indices[col_name]
            for row in ws.iter_rows(min_row=2, min_col=col_idx, max_col=col_idx):
                cell = row[0]

                cell.number_format = "General" 
    
    wb.save(output_file_local)
    wb.close()

print(f"\nArquivos gerados com sucesso: {len(locais_unicos)} planilhas diferentes.")